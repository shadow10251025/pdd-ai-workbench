"""Read merchant SKU workbooks without guessing dimension order or inventory."""
import io
from decimal import Decimal, InvalidOperation
import openpyxl
from product_package import sku_id_for_specs
from resource_limits import validate_workbook


def parse_workbook(content):
    validate_workbook(content)
    book=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        sheet=book.active
        if sheet.max_row>210 or sheet.max_column>30:
            raise ValueError('SKU表格超过200项或列数过多，请拆分后上传')
        header=None; data=[]
        for index,row in enumerate(sheet.iter_rows(values_only=True),1):
            values=[str(value).strip() if value is not None else '' for value in row]
            if index<=6 and any(value in {'库存','stock'} for value in values) and any('拼单价' in value or value in {'price','价格'} for value in values):
                header=values;continue
            if header is not None and any(value!='' for value in values): data.append(row)
        if not header:
            raise ValueError('没有找到SKU表头，请使用本工具导出的表格')
        def locate(test): return next((index for index,value in enumerate(header) if test(value)),None)
        stock_column=locate(lambda v:v in {'库存','stock'})
        price_column=locate(lambda v:'拼单价' in v or v in {'price','价格'})
        single_column=locate(lambda v:'单买价' in v or v=='single_price')
        code_column=locate(lambda v:v in {'规格编码','编码','spec_code'})
        enabled_column=locate(lambda v:v in {'启用状态','enabled','状态','是否启用'})
        name_column=locate(lambda v:v in {'规格名称','SKU名称','name'})
        reserved={stock_column,price_column,single_column,code_column,enabled_column,name_column}
        dimensions=[(index,name) for index,name in enumerate(header) if name and index not in reserved]
        def cell(row,column): return row[column] if column is not None and column<len(row) else None
        def number(value):
            if value in (None,''): return None
            try: result=Decimal(str(value))
            except (InvalidOperation,ValueError) as exc: raise ValueError('价格或库存包含无效数字') from exc
            if not result.is_finite(): raise ValueError('价格或库存必须是有限数字')
            return result
        results=[]
        for row in data:
            specs={name:str(cell(row,index)).strip() for index,name in dimensions if cell(row,index) not in (None,'')}
            name=' '.join(specs.values()) if specs else str(cell(row,name_column) or '').strip()
            if not name: continue
            stock=number(cell(row,stock_column))
            if stock is not None and (stock<0 or stock!=stock.to_integral_value()):
                raise ValueError('库存必须为大于或等于0的整数')
            price=number(cell(row,price_column)); single=number(cell(row,single_column))
            enabled=cell(row,enabled_column)
            code=cell(row,code_column)
            sku={'name':name,'price':float(price) if price is not None else 0,'single_price':float(single) if single is not None else 0,'stock':int(stock) if stock is not None else None,
                 'enabled':str(enabled).strip().lower() not in {'0','false','否','下架','停用'},'spec_code':str(code).strip() if code is not None else ''}
            if specs: sku.update(spec_values=specs,id=sku_id_for_specs(specs))
            results.append(sku)
        return results
    finally:
        book.close()
